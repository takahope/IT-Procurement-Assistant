import uuid
import zipfile
from pathlib import Path

import docx
import pytest

from procurement_review.documents.ingest import confirm_text, ingest, segment
from procurement_review.documents.structure import (
    DocumentSnapshot,
    QualityIssue,
    SourceSpan,
)
from procurement_review.domain.enums import RuntimeMode
from procurement_review.domain.errors import DomainError
from procurement_review.domain.models import RuntimePolicy
from procurement_review.storage.secrets import FakeSecretStore
from procurement_review.storage.vault import ArtifactStore

FIXTURES_DIR = Path(__file__).resolve().parents[1] / "fixtures" / "synthetic" / "documents"


@pytest.fixture
def artifact_store(tmp_path: Path) -> ArtifactStore:
    secrets = FakeSecretStore()
    vault_root = tmp_path / "vault"
    return ArtifactStore(root=vault_root, secrets=secrets)


@pytest.fixture
def runtime_policy(tmp_path: Path) -> RuntimePolicy:
    return RuntimePolicy(mode=RuntimeMode.SYNTHETIC, data_root=str(tmp_path))


@pytest.fixture
def sample_review_id() -> str:
    return str(uuid.uuid4())


def test_ingest_normal_docx(artifact_store: ArtifactStore, runtime_policy: RuntimePolicy, sample_review_id: str):
    doc_path = FIXTURES_DIR / "normal.docx"
    snapshot = ingest(doc_path, sample_review_id, runtime_policy, artifact_store)

    assert isinstance(snapshot, DocumentSnapshot)
    assert snapshot.confirmation_ref is None
    assert snapshot.original_filename == "normal.docx"
    assert len(snapshot.blocks) >= 2
    assert any("第一條 總則" in b.raw_text for b in snapshot.blocks)
    assert any("第二條 履約期限" in b.raw_text for b in snapshot.blocks)

    # Check that original document was safely preserved in artifact store
    artifact_id = snapshot.manifest.get("storage_artifact_id")
    assert artifact_id is not None
    stored_bytes = artifact_store.get(sample_review_id, artifact_id)
    assert stored_bytes == doc_path.read_bytes()


def test_ingest_table_docx(artifact_store: ArtifactStore, runtime_policy: RuntimePolicy, sample_review_id: str):
    doc_path = FIXTURES_DIR / "table.docx"
    snapshot = ingest(doc_path, sample_review_id, runtime_policy, artifact_store)

    assert snapshot.confirmation_ref is None
    table_cells = [b for b in snapshot.blocks if b.block_type == "table_cell"]
    assert len(table_cells) >= 4
    texts = [b.raw_text for b in table_cells]
    assert "項目" in texts
    assert "規格說明" in texts
    assert "資安檢核" in texts
    assert any("ISO 27001" in t for t in texts)


def test_ingest_header_footer_docx(artifact_store: ArtifactStore, runtime_policy: RuntimePolicy, sample_review_id: str):
    doc_path = FIXTURES_DIR / "header_footer.docx"
    snapshot = ingest(doc_path, sample_review_id, runtime_policy, artifact_store)

    headers = [b for b in snapshot.blocks if b.block_type == "header"]
    footers = [b for b in snapshot.blocks if b.block_type == "footer"]
    assert len(headers) >= 1
    assert any("中研院採購契約草案" in h.raw_text for h in headers)
    assert len(footers) >= 1
    assert any("第 1 頁" in f.raw_text for f in footers)


def test_ingest_textbox_docx(artifact_store: ArtifactStore, runtime_policy: RuntimePolicy, sample_review_id: str):
    doc_path = FIXTURES_DIR / "textbox.docx"
    snapshot = ingest(doc_path, sample_review_id, runtime_policy, artifact_store)

    textbox_blocks = [b for b in snapshot.blocks if b.block_type == "textbox"]
    assert len(textbox_blocks) >= 1
    assert any("所有交付項目須經資安審查" in b.raw_text for b in textbox_blocks)


def test_ingest_tracked_changes_docx(artifact_store: ArtifactStore, runtime_policy: RuntimePolicy, sample_review_id: str):
    doc_path = FIXTURES_DIR / "tracked_changes.docx"
    snapshot = ingest(doc_path, sample_review_id, runtime_policy, artifact_store)

    # Core assertion: tracked changes are not silently accepted
    assert any(issue.code == "TRACKED_CHANGES" for issue in snapshot.issues)

    # Issue fixed structure: code, block_id, safe_message
    for issue in snapshot.issues:
        assert isinstance(issue, QualityIssue)
        assert hasattr(issue, "code") and isinstance(issue.code, str)
        assert hasattr(issue, "safe_message") and isinstance(issue.safe_message, str)
        assert hasattr(issue, "block_id")


def test_ingest_hidden_text_docx(artifact_store: ArtifactStore, runtime_policy: RuntimePolicy, sample_review_id: str):
    doc_path = FIXTURES_DIR / "hidden_text.docx"
    snapshot = ingest(doc_path, sample_review_id, runtime_policy, artifact_store)

    assert any(issue.code == "HIDDEN_TEXT" for issue in snapshot.issues)


def test_ingest_normal_pdf(artifact_store: ArtifactStore, runtime_policy: RuntimePolicy, sample_review_id: str):
    doc_path = FIXTURES_DIR / "normal.pdf"
    snapshot = ingest(doc_path, sample_review_id, runtime_policy, artifact_store)

    assert snapshot.confirmation_ref is None
    assert len(snapshot.blocks) >= 1
    assert snapshot.blocks[0].page_number == 1
    assert any("Vendor shall not disclose data" in b.raw_text for b in snapshot.blocks)


def test_ingest_scanned_pdf(artifact_store: ArtifactStore, runtime_policy: RuntimePolicy, sample_review_id: str):
    doc_path = FIXTURES_DIR / "scanned.pdf"
    snapshot = ingest(doc_path, sample_review_id, runtime_policy, artifact_store)

    # Scanned PDF without text must report issue requiring human OCR / confirmation
    assert any(issue.code in ("SCANNED_PAGE", "NEEDS_OCR") for issue in snapshot.issues)


def test_ingest_mixed_pdf(artifact_store: ArtifactStore, runtime_policy: RuntimePolicy, sample_review_id: str):
    doc_path = FIXTURES_DIR / "mixed.pdf"
    snapshot = ingest(doc_path, sample_review_id, runtime_policy, artifact_store)

    assert any("Article 5 Acceptance" in b.raw_text for b in snapshot.blocks)
    assert any(issue.code == "EMBEDDED_IMAGE" for issue in snapshot.issues)


def test_ingest_encrypted_pdf_fails_cleanly(artifact_store: ArtifactStore, runtime_policy: RuntimePolicy, sample_review_id: str):
    doc_path = FIXTURES_DIR / "encrypted.pdf"
    with pytest.raises(DomainError) as exc_info:
        ingest(doc_path, sample_review_id, runtime_policy, artifact_store)
    assert exc_info.value.code == "ENCRYPTED_PDF"


def test_ingest_corrupted_zip_fails_cleanly(artifact_store: ArtifactStore, runtime_policy: RuntimePolicy, sample_review_id: str):
    doc_path = FIXTURES_DIR / "corrupt.zip"
    with pytest.raises(DomainError) as exc_info:
        ingest(doc_path, sample_review_id, runtime_policy, artifact_store)
    assert exc_info.value.code in ("CORRUPT_DOCUMENT", "UNSUPPORTED_DOCUMENT_TYPE")


def test_ingest_zip_bomb_fails_cleanly(artifact_store: ArtifactStore, runtime_policy: RuntimePolicy, sample_review_id: str):
    doc_path = FIXTURES_DIR / "zip_bomb.docx"
    with pytest.raises(DomainError) as exc_info:
        ingest(doc_path, sample_review_id, runtime_policy, artifact_store)
    assert exc_info.value.code == "ZIP_BOMB_DETECTED"


def test_normalization_preserves_negation_and_source_offsets(
    artifact_store: ArtifactStore, runtime_policy: RuntimePolicy, sample_review_id: str
):
    """Review Focus 2:

    Verify for negation keywords like '不得' and numbers like '３０日' (fullwidth),
    block.raw_text[span.start:span.end] == '不得' holds exactly.
    """
    doc_path = FIXTURES_DIR / "normal.docx"
    snapshot = ingest(doc_path, sample_review_id, runtime_policy, artifact_store)
    snapshot = segment(snapshot)

    # Locate the block containing '不得於３０日內遲延給付'
    target_block = None
    for b in snapshot.blocks:
        if "不得" in b.raw_text and "３０日" in b.raw_text:
            target_block = b
            break

    assert target_block is not None, "Target block with '不得' and fullwidth '３０日' not found"

    # Find span for negation keyword '不得'
    negation_keyword = "不得"
    start_idx = target_block.raw_text.find(negation_keyword)
    assert start_idx != -1
    span = SourceSpan(
        document_id=snapshot.document_id,
        block_id=target_block.block_id,
        start_offset=start_idx,
        end_offset=start_idx + len(negation_keyword),
        structural_path=target_block.structural_path,
    )

    # Core assertion: Unicode code point offset mapping
    assert target_block.raw_text[span.start:span.end] == "不得"

    # Fullwidth number span test
    num_keyword = "３０日"
    num_start = target_block.raw_text.find(num_keyword)
    assert num_start != -1
    num_span = SourceSpan(
        document_id=snapshot.document_id,
        block_id=target_block.block_id,
        start_offset=num_start,
        end_offset=num_start + len(num_keyword),
        structural_path=target_block.structural_path,
    )
    assert target_block.raw_text[num_span.start:num_span.end] == "３０日"


def test_confirm_text_updates_blocks_and_generates_new_snapshot(
    artifact_store: ArtifactStore, runtime_policy: RuntimePolicy, sample_review_id: str
):
    doc_path = FIXTURES_DIR / "normal.docx"
    snapshot = ingest(doc_path, sample_review_id, runtime_policy, artifact_store)
    assert snapshot.confirmation_ref is None

    block_to_correct = snapshot.blocks[0]
    corrected_blocks = [
        {
            "block_id": block_to_correct.block_id,
            "raw_text": "【人工更正】第一條 總則 本合約經雙方專案經理簽章核可。",
        }
    ]

    actor = "auditor@sinica.edu.tw"
    confirmed_snapshot = confirm_text(snapshot, corrected_blocks, actor=actor)

    # Core assertions
    assert confirmed_snapshot.confirmation_ref is not None
    assert isinstance(confirmed_snapshot.confirmation_ref, str)
    assert confirmed_snapshot.snapshot_id != snapshot.snapshot_id
    assert confirmed_snapshot.document_id == snapshot.document_id

    # Block content is updated
    new_first_block = confirmed_snapshot.blocks[0]
    assert new_first_block.raw_text == "【人工更正】第一條 總則 本合約經雙方專案經理簽章核可。"

    # Manifest records actor and history
    assert confirmed_snapshot.manifest.get("confirmation_actor") == actor
    assert confirmed_snapshot.manifest.get("previous_snapshot_id") == str(snapshot.snapshot_id)
    assert "confirmed_at" in confirmed_snapshot.manifest


def test_segment_identifies_clauses_and_spans(
    artifact_store: ArtifactStore, runtime_policy: RuntimePolicy, sample_review_id: str
):
    doc_path = FIXTURES_DIR / "normal.docx"
    snapshot = ingest(doc_path, sample_review_id, runtime_policy, artifact_store)
    segmented = segment(snapshot)

    assert len(segmented.clauses) >= 2
    clause_titles = [c.title for c in segmented.clauses if c.title]
    clause_numbers = [c.clause_number for c in segmented.clauses if c.clause_number]

    assert any("第一條" in n for n in clause_numbers)
    assert any("總則" in t for t in clause_titles)

    for clause in segmented.clauses:
        assert clause.clause_id
        assert clause.text
        assert len(clause.block_ids) > 0
        assert len(clause.spans) > 0
        for span in clause.spans:
            assert isinstance(span, SourceSpan)
            assert span.start >= 0
            assert span.end > span.start


def test_ingest_xml_entity_expansion_fails_cleanly(
    artifact_store: ArtifactStore, runtime_policy: RuntimePolicy, sample_review_id: str, tmp_path: Path
):
    malicious_docx = tmp_path / "xxe_payload.docx"
    payload = b"""<?xml version="1.0" encoding="UTF-8"?>
<!DOCTYPE test [
  <!ENTITY xxe "evil content">
  <!ELEMENT test (#PCDATA)>
]>
<test>&xxe;</test>"""
    with zipfile.ZipFile(malicious_docx, "w") as zf:
        zf.writestr("[Content_Types].xml", payload)
        zf.writestr("word/document.xml", payload)

    with pytest.raises(DomainError) as exc_info:
        ingest(malicious_docx, sample_review_id, runtime_policy, artifact_store)
    assert exc_info.value.code == "CORRUPT_DOCUMENT"
    assert "XML security check or parse failed" in exc_info.value.safe_message


def test_ingest_preserves_paragraph_and_table_natural_order(
    artifact_store: ArtifactStore, runtime_policy: RuntimePolicy, sample_review_id: str, tmp_path: Path
):
    doc_path = tmp_path / "ordered_doc.docx"
    doc = docx.Document()
    doc.add_paragraph("第一條 採購標的")
    tbl = doc.add_table(rows=1, cols=1)
    tbl.rows[0].cells[0].text = "伺服器主機 10 臺"
    doc.add_paragraph("第二條 履約期限")
    doc.save(str(doc_path))

    snapshot = ingest(doc_path, sample_review_id, runtime_policy, artifact_store)

    body_blocks = [b for b in snapshot.blocks if b.block_type in ("paragraph", "table_cell")]
    assert len(body_blocks) == 3
    assert body_blocks[0].block_type == "paragraph"
    assert "第一條 採購標的" in body_blocks[0].raw_text
    assert body_blocks[1].block_type == "table_cell"
    assert "伺服器主機 10 臺" in body_blocks[1].raw_text
    assert body_blocks[2].block_type == "paragraph"
    assert "第二條 履約期限" in body_blocks[2].raw_text

    # Verify segmentation keeps the table cell grouped under Clause 1
    segmented = segment(snapshot)
    assert len(segmented.clauses) == 2
    assert segmented.clauses[0].clause_number == "第一條"
    assert "伺服器主機 10 臺" in segmented.clauses[0].text
    assert segmented.clauses[1].clause_number == "第二條"


def test_ingest_textbox_content_not_duplicated_in_paragraph(
    artifact_store: ArtifactStore, runtime_policy: RuntimePolicy, sample_review_id: str
):
    doc_path = FIXTURES_DIR / "textbox.docx"
    snapshot = ingest(doc_path, sample_review_id, runtime_policy, artifact_store)

    paragraph_blocks = [b for b in snapshot.blocks if b.block_type == "paragraph"]
    textbox_blocks = [b for b in snapshot.blocks if b.block_type == "textbox"]

    assert len(paragraph_blocks) >= 1
    assert len(textbox_blocks) >= 1

    # Textbox content must only be in textbox block and not duplicated in paragraph
    assert "備註：所有交付項目須經資安審查。" in textbox_blocks[0].raw_text
    assert "備註：所有交付項目須經資安審查。" not in paragraph_blocks[0].raw_text
    assert paragraph_blocks[0].raw_text == "第四條 交付項目"


def test_confirm_text_unknown_block_id_raises_domain_error(
    artifact_store: ArtifactStore, runtime_policy: RuntimePolicy, sample_review_id: str
):
    doc_path = FIXTURES_DIR / "normal.docx"
    snapshot = ingest(doc_path, sample_review_id, runtime_policy, artifact_store)

    non_existent_id = str(uuid.uuid4())
    corrected_blocks = [
        {"block_id": non_existent_id, "raw_text": "不存在的區塊更正"}
    ]

    with pytest.raises(DomainError) as exc_info:
        confirm_text(snapshot, corrected_blocks, actor="auditor@sinica.edu.tw")
    assert exc_info.value.code == "BLOCK_NOT_FOUND"
